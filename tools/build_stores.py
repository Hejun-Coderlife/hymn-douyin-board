# -*- coding: utf-8 -*-
"""从抖音来客导出的 xlsx 重建 data/stores.json。

用法: python3 tools/build_stores.py [xlsx路径]
不传路径就用 reference/ 下的样例。
已有的 stores.json 里手工填过的字段（区域/门店类型/...）会被保留。
"""
import json
import pathlib
import re
import sys

import openpyxl

ROOT = pathlib.Path(__file__).resolve().parent.parent
OUT = ROOT / "data" / "stores.json"
BRAND_ID = "74892160771"  # 品牌号「赫眉美妆护肤连锁」，不进看板

# 预留字段，留空等后续人工补
BLANK = {
    "region": "",
    "storeType": "",
    "customer": "",
    "mainService": "",
    "scenes": [],
    "onCamera": "",
    "startDate": None,
    "note": "",
}

# 新账号：三横店 8/24 才开始发视频
START_DATES = {"28098977805": "2026-08-24"}


def parse_name(account_name):
    """「赫眉·珀莱雅(新河店)」-> ('新河店', '珀莱雅')；「赫眉(西坞店)」-> ('西坞店', '')"""
    name = (account_name or "").strip()
    m = re.search(r"[（(]\s*(.+?)\s*[）)]", name)
    store = m.group(1) if m else name.replace("赫眉", "").strip("·- ")
    head = name[: m.start()] if m else name
    line = head.replace("赫眉", "").strip("·- ")
    return store, line


def main():
    src = sys.argv[1] if len(sys.argv) > 1 else str(
        ROOT / "reference" / "生意经_视频列表_20260916_1342.xlsx"
    )
    ws = openpyxl.load_workbook(src)["Sheet1"]
    header = [c.value for c in next(ws.iter_rows(max_row=1))]
    i_id, i_name = header.index("抖音号"), header.index("抖音号名称")

    seen = {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        did = str(row[i_id]).strip()
        if did == BRAND_ID or did in seen:
            continue
        seen[did] = str(row[i_name] or "").strip()

    old = {}
    if OUT.exists():
        old = {s["douyinId"]: s for s in json.loads(OUT.read_text(encoding="utf-8"))}

    stores = []
    for did, account in seen.items():
        store_name, brand_line = parse_name(account)
        rec = {"douyinId": did, "accountName": account,
               "storeName": store_name, "brandLine": brand_line}
        rec.update(BLANK)
        if did in START_DATES:
            rec["startDate"] = START_DATES[did]
            rec["note"] = "新账号，2026-08-24 才开始发视频"
        # 保留人工填过的内容
        for k, v in old.get(did, {}).items():
            if k in BLANK and v not in ("", [], None):
                rec[k] = v
        stores.append(rec)

    stores.sort(key=lambda s: s["storeName"])

    # 门店编号：S01~S43，仅用于脚本 id 的可读性；已有编号保持不变，新店顺延
    used = {o.get("code") for o in old.values() if o.get("code")}
    nxt = 1
    for rec in stores:
        keep = old.get(rec["douyinId"], {}).get("code")
        if keep:
            rec["code"] = keep
            continue
        while f"S{nxt:02d}" in used:
            nxt += 1
        rec["code"] = f"S{nxt:02d}"
        used.add(rec["code"])
    for rec in stores:  # code 排到前面，便于阅读
        rec_items = [("code", rec.pop("code"))] + list(rec.items())
        rec.clear(); rec.update(dict(rec_items))
    OUT.write_text(json.dumps(stores, ensure_ascii=False, indent=2), encoding="utf-8")
    print(f"写出 {OUT} — {len(stores)} 家门店（已排除品牌号 {BRAND_ID}）")


if __name__ == "__main__":
    main()
